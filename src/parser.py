import openpyxl
from src.config import EXCEL_PATH, DAYS_CONFIG

def clean_dummy_markers(text: str) -> str:
    """Убирает одиночные крестики (x / х), прочерки и точки, превращая их в пустое окно"""
    if not text:
        return ""
    cleaned = text.strip()
    # Убираем русский 'х', английский 'x', тире, прочерки и точки
    if cleaned.lower() in ["x", "х", "-", "—", ".", "окно", "нет"]:
        return ""
    return cleaned

def get_cell_value(sheet, row: int, col: int) -> str:
    """Считывает текст ячейки с учетом объединения и очищает от крестиков-заглушек"""
    val = sheet.cell(row=row, column=col).value
    if val is not None:
        return clean_dummy_markers(str(val))

    # Если ячейка объединена с соседней
    for rng in sheet.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            top_val = sheet.cell(row=rng.min_row, column=rng.min_col).value
            return clean_dummy_markers(str(top_val)) if top_val is not None else ""
    return ""

def get_available_groups() -> list[str]:
    """Считывает список всех доступных групп из первой строки"""
    if not EXCEL_PATH.exists():
        return []
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb.worksheets[0]
    groups = []
    for col in range(4, sheet.max_column + 1):
        val = sheet.cell(row=1, column=col).value
        if val:
            g = str(val).strip()
            if g and g not in groups:
                groups.append(g)
    return groups

def parse_week_schedule(group_name: str, week_num: int):
    """Парсит расписание на выбранную неделю"""
    if not EXCEL_PATH.exists():
        return None

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet_name = f"1курс-{week_num}неделя"
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    sheet = wb[sheet_name]

    group_col = None
    for col in range(4, sheet.max_column + 1):
        val = sheet.cell(row=1, column=col).value
        if val and str(val).strip().lower() == group_name.strip().lower():
            group_col = col
            break

    if not group_col:
        return None

    schedule_data = []
    for day_full, day_short, start_row in DAYS_CONFIG:
        pairs = {}
        for p in range(1, 9):
            row = start_row + (p - 1)
            pairs[p] = get_cell_value(sheet, row, group_col)

        schedule_data.append({
            "name": day_short,
            "pairs": pairs
        })

    return schedule_data